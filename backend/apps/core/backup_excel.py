"""Build NHMS backup as one Excel workbook (multiple sheets) inside a zip."""

from __future__ import annotations

import json
from io import BytesIO
import zipfile

from openpyxl import Workbook


def _excel_cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (list, dict)):
        return json.dumps(value, default=str)
    return str(value)


def _write_rows_to_sheet(sheet, rows: list[dict]) -> None:
    if not rows:
        sheet.append(["(no rows)"])
        return
    headers: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                headers.append(key)
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_cell_value(row.get(header)) for header in headers])


def build_excel_workbook(payload: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    skip = {"format", "generatedAt"}
    for key, value in payload.items():
        if key in skip or not isinstance(value, list):
            continue
        sheet = workbook.create_sheet(title=key[:31])
        _write_rows_to_sheet(sheet, value)
    if not workbook.sheetnames:
        sheet = workbook.create_sheet(title="info")
        sheet.append(["(no data)"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_excel_zip(payload: dict) -> bytes:
    xlsx = build_excel_workbook(payload)
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "README.txt",
            "NHMS daily backup.\n"
            "Open nhms-backup.xlsx in Microsoft Excel or LibreOffice.\n"
            "Each worksheet tab is one data table (users, patients, stock, etc.).\n",
        )
        archive.writestr("nhms-backup.xlsx", xlsx)
    return buffer.getvalue()
